"""出入库流水：查询、导出、模板下载、批量导入。"""
from __future__ import annotations

import csv
import io
import sqlite3

from flask import Blueprint, Response, jsonify, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from ..db import get_db, to_dicts
from ..security import current_user, superadmin_required
from ..services import importer
from ..services.importer import ImportError_
from ..services.inventory import TX_IMPORT, log_transaction

bp = Blueprint("records", __name__, url_prefix="/api/records")

RECORD_SQL = """
SELECT t.id, t.component_id, t.delta, t.type, t.operator, t.remark, t.applicant, t.created_at,
       IFNULL(c.name, '已删除') AS name, IFNULL(c.spec, '') AS spec, IFNULL(c.unit, '') AS unit
FROM transactions t
LEFT JOIN components c ON c.id = t.component_id
{where}
ORDER BY t.id DESC
LIMIT ? OFFSET ?
"""

COUNT_SQL = """
SELECT COUNT(*) AS total FROM transactions t
LEFT JOIN components c ON c.id = t.component_id
{where}
"""

EXPORT_HEADER = ["设备名称", "设备固定资产编号", "分类", "设备型号", "单位", "位置", "设备负责人", "设备状态", "备注"]
TEMPLATE_ROWS = [
    ["巡检机器人", "ZC-2026-001", "移动机器人", "RB-01", "台", "特种环境实验室A区", "张老师", "可借用", "每行代表一台设备"],
    ["防爆机械臂", "ZC-2026-002", "作业设备", "ARM-EX", "台", "特种环境实验室B区", "李老师", "维护中", "每行代表一台设备"],
]


def _conditions():
    clauses, args = [], []
    component_id = request.args.get("cid")
    if component_id:
        clauses.append("t.component_id = ?")
        args.append(component_id)

    kind = (request.args.get("type") or "").strip()
    if kind:
        clauses.append("t.type = ?")
        args.append(kind)

    keyword = (request.args.get("q") or "").strip()
    if keyword:
        # 关键词覆盖器件名、规格、操作人、申请人与备注，服务端过滤后分页
        clauses.append(
            "(c.name LIKE ? OR c.spec LIKE ? OR t.operator LIKE ? "
            "OR t.applicant LIKE ? OR t.remark LIKE ?)"
        )
        args.extend([f"%{keyword}%"] * 5)

    return ("WHERE " + " AND ".join(clauses) if clauses else ""), args


@bp.route("")
@superadmin_required
def list_records():
    db = get_db()
    where, args = _conditions()
    limit = min(max(request.args.get("limit", 50, type=int) or 50, 1), 500)
    offset = max(request.args.get("offset", 0, type=int) or 0, 0)

    rows = db.execute(RECORD_SQL.format(where=where), [*args, limit, offset]).fetchall()
    total = db.execute(COUNT_SQL.format(where=where), args).fetchone()["total"]
    return jsonify({"ok": True, "data": to_dicts(rows), "total": total})


def _csv_response(text, filename):
    return Response(
        "\ufeff" + text,  # BOM，保证 Excel 正确识别中文
        mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


def _to_csv(rows):
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerows(rows)
    return buffer.getvalue()


@bp.route("/export")
@superadmin_required
def export_components():
    db = get_db()
    rows = db.execute(
        "WITH listed AS (SELECT c.*, "
        "(SELECT ROUND(IFNULL(SUM(p.qty), 0), 6) FROM custody_positions p "
        " WHERE p.component_id = c.id AND ROUND(p.qty, 6) > 0) AS borrowed_qty "
        "FROM components c) "
        "SELECT name, asset_no, category, spec, unit, location, owner, "
        "CASE WHEN stock > 0 THEN '可借用' WHEN borrowed_qty > 0 THEN '借出中' "
        "WHEN maintenance_qty > 0 THEN '维护中' ELSE '不可借' END AS equipment_status, "
        "remark FROM listed ORDER BY category, name, asset_no"
    ).fetchall()
    body = _to_csv([EXPORT_HEADER, *([row[key] for key in row.keys()] for row in rows)])
    return _csv_response(body, "components.csv")


@bp.route("/template")
@superadmin_required
def download_template():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "设备信息"
    sheet.append(EXPORT_HEADER)
    for row in TEMPLATE_ROWS:
        sheet.append(row)

    header_fill = PatternFill("solid", fgColor="2F7FD3")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10, color="243447")
    border = Border(bottom=Side(style="thin", color="D5DFEA"))
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="center")
            cell.border = border

    widths = [20, 22, 16, 18, 10, 24, 16, 14, 28]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.row_dimensions[1].height = 24
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:I{sheet.max_row}"
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["B"].number_format = "@"
    status_validation = DataValidation(
        type="list",
        formula1='"可借用,维护中,不可借"',
        allow_blank=True,
    )
    status_validation.error = "请选择可借用、维护中或不可借"
    status_validation.errorTitle = "设备状态无效"
    sheet.add_data_validation(status_validation)
    status_validation.add("H2:H1000")

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name="设备批量导入模板.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@bp.route("/import", methods=["POST"])
@superadmin_required
def import_file():
    upload = request.files.get("file")
    if upload is None:
        return jsonify({"ok": False, "msg": "未收到上传文件"}), 400

    try:
        items, skipped = importer.parse(upload.read(), upload.filename or "")
    except ImportError_ as exc:
        return jsonify({"ok": False, "msg": str(exc)}), 400
    except Exception as exc:  # 解析库抛出的底层异常
        return jsonify({"ok": False, "msg": f"解析失败: {exc}"}), 400

    if not items:
        return jsonify({"ok": False, "msg": "没有可导入的有效数据行"}), 400

    db = get_db()
    operator = current_user()
    imported = 0
    try:
        for item in items:
            cursor = db.execute(
                "INSERT INTO components "
                "(name, asset_no, category, spec, unit, location, owner, maintenance, unavailable, "
                "maintenance_qty, unavailable_qty, stock, threshold, remark) "
                "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0, ?)",
                (
                    item["name"], item["asset_no"], item["category"], item["spec"],
                    item["unit"], item["location"], item["owner"],
                    item["maintenance"], item["unavailable"],
                    1 if item["maintenance"] else 0,
                    1 if item["unavailable"] else 0,
                    0 if item["maintenance"] or item["unavailable"] else 1,
                    item["remark"],
                ),
            )
            log_transaction(db, cursor.lastrowid, 1, TX_IMPORT, operator, "导入单台设备")
            imported += 1
        db.commit()
    except sqlite3.IntegrityError:
        db.rollback()
        return jsonify({"ok": False, "msg": "固定资产编号重复，请确保每台设备编号唯一"}), 400
    except Exception as exc:
        db.rollback()
        return jsonify({"ok": False, "msg": f"写入失败: {exc}"}), 500

    return jsonify({"ok": True, "imported": imported, "skipped": skipped})
