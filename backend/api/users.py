"""普通用户账号管理接口（仅管理员）。"""
from __future__ import annotations

import sqlite3

import io

from flask import Blueprint, jsonify, request, send_file, session
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from ..config import Config
from ..db import get_db, to_dicts
from ..security import manager_required, superadmin_required
from ..services import user_importer
from ..services import admin_profile
from ..services import users as user_service
from ..services.user_importer import UserImportError
from ..services.users import UserError

bp = Blueprint("users", __name__, url_prefix="/api/users")


@bp.errorhandler(UserError)
def on_user_error(exc):
    return jsonify({"ok": False, "msg": str(exc)}), 400


@bp.route("", methods=["GET"])
@superadmin_required
def list_users():
    db = get_db()
    rows = user_service.list_users(db)
    return jsonify({"ok": True, "data": to_dicts(rows)})


@bp.route("/superadmin", methods=["GET", "PUT"])
@superadmin_required
def superadmin_profile():
    db = get_db()
    if request.method == "GET":
        return jsonify({"ok": True, "data": admin_profile.get_profile(db)})

    payload = request.get_json(silent=True) or {}
    try:
        profile = admin_profile.update_profile(
            db, payload.get("full_name"), payload.get("phone")
        )
        db.commit()
    except admin_profile.AdminProfileError as exc:
        db.rollback()
        return jsonify({"ok": False, "msg": str(exc)}), 400

    session["user"] = profile["username"]
    session["full_name"] = profile["full_name"]
    session["phone"] = profile["phone"]
    return jsonify({"ok": True, "data": profile, "msg": "总管理员信息已保存"})


@bp.route("/directory", methods=["GET"])
@manager_required
def borrower_directory():
    """借还办理使用的精简通讯录，不返回角色、密码等管理信息。"""
    rows = get_db().execute(
        "SELECT username, full_name, phone FROM users "
        "WHERE phone <> '' ORDER BY full_name, username"
    ).fetchall()
    data = [dict(row) for row in rows]
    profile = admin_profile.get_profile(get_db())
    if profile["phone"]:
        data.insert(0, {
            "username": profile["username"],
            "full_name": profile["full_name"],
            "phone": profile["phone"],
        })
    return jsonify({"ok": True, "data": data})


@bp.route("", methods=["POST"])
@superadmin_required
def create_user():
    payload = request.get_json(silent=True) or {}
    username = (payload.get("username") or "").strip()
    if username == Config.USERNAME:
        return jsonify({"ok": False, "msg": "该用户名是内置管理员账号，不能重复使用"}), 400

    db = get_db()
    try:
        user_id = user_service.create_user(
            db,
            username,
            role=payload.get("role", "member"),
            user_type=payload.get("user_type"),
            mentor=payload.get("mentor", ""),
            full_name=payload.get("full_name", ""),
            phone=payload.get("phone", ""),
            reserved_names=(Config.USERNAME,),
        )
        db.commit()
    except UserError as exc:
        db.rollback()
        return jsonify({"ok": False, "msg": str(exc)}), 400
    except sqlite3.IntegrityError:
        db.rollback()
        return jsonify({"ok": False, "msg": "用户名已存在"}), 400

    return jsonify(
        {
            "ok": True,
            "id": user_id,
            "initial_password": user_service.default_password(username),
            "msg": f"已创建用户 {username}",
        }
    ), 201


@bp.route("/import", methods=["POST"])
@superadmin_required
def import_users():
    upload = request.files.get("file")
    if upload is None:
        return jsonify({"ok": False, "msg": "请选择 Excel 文件"}), 400
    try:
        items = user_importer.parse(upload.read(), upload.filename or "")
    except UserImportError as exc:
        return jsonify({"ok": False, "msg": str(exc)}), 400

    db = get_db()
    imported = 0
    errors = []
    try:
        for item in items:
            if item["username"] == Config.USERNAME:
                errors.append(f"第 {item['row']} 行：账号为总管理员保留账号")
                continue
            try:
                user_service.create_user(
                    db,
                    item["username"],
                    role=item["role"],
                    user_type=item["user_type"],
                    mentor=item["mentor"],
                    full_name=item["full_name"],
                    phone=item["phone"],
                    reserved_names=(Config.USERNAME,),
                )
                imported += 1
            except (UserError, sqlite3.IntegrityError) as exc:
                message = str(exc) if str(exc) else "账号已存在"
                errors.append(f"第 {item['row']} 行：{message}")
        if not imported:
            db.rollback()
            return jsonify({"ok": False, "msg": errors[0] if errors else "没有可导入的用户"}), 400
        db.commit()
    except Exception as exc:
        db.rollback()
        return jsonify({"ok": False, "msg": f"导入失败：{exc}"}), 500

    return jsonify(
        {
            "ok": True,
            "imported": imported,
            "skipped": len(errors),
            "errors": errors[:10],
            "msg": f"成功导入 {imported} 个用户" + (f"，跳过 {len(errors)} 行" if errors else ""),
        }
    )


@bp.route("/template", methods=["GET"])
@superadmin_required
def download_user_template():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "用户信息"
    sheet.append(["用户类型", "姓名", "账号", "手机号", "所属导师", "权限角色"])
    sheet.append(["老师", "张老师", "13800000000", "13800000000", "", "管理员"])
    sheet.append(["学生", "李同学", "20260001", "13900000000", "张老师", "普通用户"])

    header_fill = PatternFill("solid", fgColor="2F7FD3")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10, color="243447")
    border = Border(bottom=Side(style="thin", color="D5DFEA"))
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="center")
            cell.border = border

    for column, width in zip(("A", "B", "C", "D", "E", "F"), (16, 18, 22, 20, 20, 16)):
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 24
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:F3"
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["C"].number_format = "@"

    type_validation = DataValidation(type="list", formula1='"老师,学生"', allow_blank=False)
    role_validation = DataValidation(type="list", formula1='"普通用户,管理员"', allow_blank=True)
    sheet.add_data_validation(type_validation)
    sheet.add_data_validation(role_validation)
    type_validation.add("A2:A1000")
    role_validation.add("F2:F1000")

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name="用户批量导入模板.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@bp.route("/<int:user_id>/password", methods=["PUT"])
@superadmin_required
def reset_password(user_id):
    payload = request.get_json(silent=True) or {}
    db = get_db()
    try:
        username, _version = user_service.set_password(db, user_id, payload.get("password"))
        db.commit()
    except UserError as exc:
        db.rollback()
        return jsonify({"ok": False, "msg": str(exc)}), 400
    return jsonify({"ok": True, "msg": f"已重置 {username} 的密码，该账号旧会话已失效"})


@bp.route("/<int:user_id>/role", methods=["PUT"])
@superadmin_required
def change_role(user_id):
    payload = request.get_json(silent=True) or {}
    db = get_db()
    try:
        username, role = user_service.set_role(db, user_id, payload.get("role"))
        db.commit()
    except UserError as exc:
        db.rollback()
        return jsonify({"ok": False, "msg": str(exc)}), 400
    label = "管理员" if role == "manager" else "普通用户"
    return jsonify({"ok": True, "msg": f"已将 {username} 设为{label}，其旧会话已失效"})


@bp.route("/<int:user_id>/mentor", methods=["PUT"])
@superadmin_required
def change_mentor(user_id):
    payload = request.get_json(silent=True) or {}
    db = get_db()
    try:
        username, mentor = user_service.set_mentor(db, user_id, payload.get("mentor"))
        db.commit()
    except UserError as exc:
        db.rollback()
        return jsonify({"ok": False, "msg": str(exc)}), 400
    return jsonify({"ok": True, "mentor": mentor, "msg": f"已更新 {username} 的所属导师"})


@bp.route("/<int:user_id>/name", methods=["PUT"])
@superadmin_required
def change_name(user_id):
    payload = request.get_json(silent=True) or {}
    db = get_db()
    try:
        username, full_name = user_service.set_full_name(
            db, user_id, payload.get("full_name")
        )
        db.commit()
    except UserError as exc:
        db.rollback()
        return jsonify({"ok": False, "msg": str(exc)}), 400
    return jsonify(
        {
            "ok": True,
            "full_name": full_name,
            "msg": f"已更新 {username} 的姓名",
        }
    )


@bp.route("/<int:user_id>/phone", methods=["PUT"])
@superadmin_required
def change_phone(user_id):
    payload = request.get_json(silent=True) or {}
    db = get_db()
    try:
        username, phone = user_service.set_phone(db, user_id, payload.get("phone"))
        db.commit()
    except UserError as exc:
        db.rollback()
        return jsonify({"ok": False, "msg": str(exc)}), 400
    return jsonify({"ok": True, "phone": phone, "msg": f"已更新 {username} 的手机号"})


@bp.route("/<int:user_id>", methods=["DELETE"])
@superadmin_required
def delete_user(user_id):
    db = get_db()
    try:
        username = user_service.delete_user(db, user_id)
        db.commit()
    except UserError as exc:
        db.rollback()
        return jsonify({"ok": False, "msg": str(exc)}), 400
    return jsonify({"ok": True, "msg": f"已删除用户 {username}，其旧会话已失效"})
