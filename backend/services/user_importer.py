"""用户 Excel 批量导入解析。

只负责把工作簿转换成标准行；账号格式、重复账号和密码规则由 users 服务统一校验。
"""
from __future__ import annotations

import io

from openpyxl import load_workbook


SUPPORTED_SUFFIXES = (".xlsx", ".xlsm")
MAX_ROWS = 1000

COLUMN_ALIASES = {
    "user_type": ("用户类型", "身份", "类型", "user_type", "type"),
    "full_name": ("姓名", "用户姓名", "真实姓名", "name", "full_name"),
    "username": ("账号", "用户名", "手机号", "学号", "username", "account"),
    "phone": ("手机号", "联系电话", "手机", "phone", "mobile"),
    "mentor": ("所属导师", "导师", "指导老师", "mentor"),
    "role": ("权限角色", "角色", "权限", "role"),
}

TYPE_ALIASES = {
    "老师": "teacher",
    "教师": "teacher",
    "teacher": "teacher",
    "学生": "student",
    "student": "student",
}

ROLE_ALIASES = {
    "": "member",
    "普通用户": "member",
    "用户": "member",
    "member": "member",
    "管理员": "manager",
    "manager": "manager",
}


class UserImportError(Exception):
    """可直接返回给界面的用户导入错误。"""


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _column_map(header):
    lookup = {
        alias.lower(): field
        for field, aliases in COLUMN_ALIASES.items()
        for alias in aliases
    }
    result = {}
    for index, value in enumerate(header):
        field = lookup.get(_text(value).lower())
        if field and field not in result:
            result[field] = index
    return result


def _cell(row, columns, field):
    index = columns.get(field)
    return _text(row[index]) if index is not None and index < len(row) else ""


def parse(raw, filename):
    suffix = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix not in SUPPORTED_SUFFIXES:
        raise UserImportError("仅支持 .xlsx / .xlsm Excel 文件")
    if not raw:
        raise UserImportError("文件内容为空")

    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise UserImportError(f"无法读取 Excel 文件：{exc}") from exc

    try:
        rows = [[cell.value for cell in row] for row in workbook.active.iter_rows()]
    finally:
        workbook.close()

    if not rows:
        raise UserImportError("Excel 文件没有数据")
    columns = _column_map(rows[0])
    missing = [
        label
        for field, label in (
            ("user_type", "用户类型"),
            ("full_name", "姓名"),
            ("username", "账号"),
            ("phone", "手机号"),
        )
        if field not in columns
    ]
    if missing:
        raise UserImportError(f"Excel 缺少必填列：{'、'.join(missing)}")

    items = []
    for row_number, row in enumerate(rows[1 : MAX_ROWS + 1], start=2):
        if not row or all(_text(value) == "" for value in row):
            continue
        raw_type = _cell(row, columns, "user_type").lower()
        raw_role = _cell(row, columns, "role").lower()
        items.append(
            {
                "row": row_number,
                "full_name": _cell(row, columns, "full_name"),
                "username": _cell(row, columns, "username"),
                "phone": _cell(row, columns, "phone"),
                "user_type": TYPE_ALIASES.get(raw_type, raw_type),
                "mentor": _cell(row, columns, "mentor"),
                "role": ROLE_ALIASES.get(raw_role, raw_role),
            }
        )

    if len(rows) > MAX_ROWS + 1:
        raise UserImportError(f"单次最多导入 {MAX_ROWS} 个用户")
    if not items:
        raise UserImportError("Excel 中没有可导入的用户数据")
    return items
