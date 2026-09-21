"""批量导入：把 Excel / CSV 解析成规范化的器件列表。

分两段：解析（纯函数，不占用数据库连接）+ 落库（由调用方开短事务）。
这样即使上传大文件，也不会长时间持有 SQLite 写锁。
"""
from __future__ import annotations

import csv
import io

from openpyxl import load_workbook

from .catalog import normalize, to_number

# 表头别名 -> 内部字段名。列名匹配大小写不敏感，兼容中英文表头
COLUMN_ALIASES = {
    "name": ("name", "名称", "器件", "器件名称", "物料名称", "品名"),
    "category": ("category", "分类", "类别", "类型"),
    "spec": ("spec", "规格", "型号", "规格型号"),
    "unit": ("unit", "单位"),
    "location": ("location", "位置", "存放位置", "库位", "货位"),
    "stock": ("stock", "数量", "库存", "当前数量", "当前库存"),
    "threshold": ("threshold", "阈值", "预警值", "下限", "安全库存"),
    "remark": ("remark", "备注", "说明"),
}

SUPPORTED_SUFFIXES = (".csv", ".xlsx", ".xlsm")


class ImportError_(Exception):
    """导入过程中的可预期错误，消息直接面向用户。"""


def build_column_map(header):
    """把表头行映射成 {字段名: 列下标}。"""
    lookup = {}
    for field, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            lookup[alias] = field

    column_map = {}
    for index, cell in enumerate(header):
        key = ("" if cell is None else str(cell)).strip().lower()
        field = lookup.get(key)
        # 同一字段只取第一次出现的列
        if field and field not in column_map:
            column_map[field] = index
    return column_map


def _decode(raw):
    for encoding in ("utf-8-sig", "gbk"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def read_rows(raw, filename):
    """把文件字节读成二维列表；不认识的后缀直接报错。"""
    suffix = ("." + (filename or "").rsplit(".", 1)[-1].lower()) if "." in (filename or "") else ""
    if suffix not in SUPPORTED_SUFFIXES:
        raise ImportError_("仅支持 .csv / .xlsx / .xlsm 文件")

    if suffix == ".csv":
        return list(csv.reader(io.StringIO(_decode(raw))))

    workbook = load_workbook(filename=io.BytesIO(raw), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        return [[cell.value for cell in row] for row in sheet.iter_rows()]
    finally:
        workbook.close()


def _cell(row, column_map, field, default=""):
    index = column_map.get(field)
    if index is None or index >= len(row):
        return default
    value = row[index]
    return default if value is None else str(value).strip()


def extract_items(rows):
    """二维表 -> (规范化器件列表, 跳过行数)。"""
    if not rows:
        raise ImportError_("文件内容为空")

    column_map = build_column_map(rows[0])
    has_header = bool(column_map)
    # 认不出表头时，退化为「第一列即名称」
    if not has_header:
        column_map = {"name": 0}

    items, skipped = [], 0
    for row in rows[1:] if has_header else rows:
        if not row or all(value is None or str(value).strip() == "" for value in row):
            continue
        fields = normalize(
            {
                "name": _cell(row, column_map, "name"),
                "category": _cell(row, column_map, "category"),
                "spec": _cell(row, column_map, "spec"),
                "unit": _cell(row, column_map, "unit"),
                "location": _cell(row, column_map, "location"),
                "threshold": _cell(row, column_map, "threshold", "0"),
                "remark": _cell(row, column_map, "remark"),
            }
        )
        if not fields["name"]:
            skipped += 1
            continue
        fields["stock"] = to_number(_cell(row, column_map, "stock", "0"))
        items.append(fields)

    return items, skipped


def parse(raw, filename):
    """完整解析流程：字节 -> (器件列表, 跳过行数)。"""
    return extract_items(read_rows(raw, filename))