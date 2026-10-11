"""后端接口冒烟测试：登录 -> 建器件 -> 出入库 -> 流转精度 -> 用户权限 -> 跨页检索。

用法：python backend/tests/smoke.py  （使用临时数据库，跑完自动清理）

覆盖重点（对应交接文档的回归要求）：
1. 0.3 分次转交 0.1、0.2 后无残量；分次归还库存准确；过多小数位与超量操作被拒绝
2. 跨页检索：目标记录在未筛选的后续页，搜索仍可返回；历史持有人可检索；总数与分页一致
3. 普通用户权限边界、会话失效；出库申请人必填与检索；有流转历史的器件禁止删除
4. 登录限流双维度：账号锁定（防换 IP 定向爆破）、IP 锁定（防喷洒）、互不牵连
"""
import hashlib
import hmac
import io
import os
import shutil
import sqlite3
import sys
import tempfile
import time
from datetime import date, timedelta

from openpyxl import Workbook, load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

TMP_DB = os.path.join(tempfile.gettempdir(), "components_smoke.db")
for suffix in ("", "-wal", "-shm"):
    if os.path.exists(TMP_DB + suffix):
        os.remove(TMP_DB + suffix)

# 部署脚本必须指向临时脚本：否则签名合法的那个用例会真的执行仓库里的
# deploy.sh，跑起 git pull 和 systemctl restart。
SMOKE_SECRET = "smoke-webhook-secret"
SMOKE_SCRIPT = os.path.join(tempfile.gettempdir(), "components_smoke_deploy.sh")
SMOKE_MARKER = os.path.join(tempfile.gettempdir(), "components_smoke_deploy.marker")
SMOKE_LOG = os.path.join(tempfile.gettempdir(), "components_smoke_deploy.log")

os.environ["APP_DB_PATH"] = TMP_DB
os.environ["WEBHOOK_SECRET"] = SMOKE_SECRET
os.environ["APP_DEPLOY_SCRIPT"] = SMOKE_SCRIPT
os.environ["APP_DEPLOY_LOG"] = SMOKE_LOG

from flask.sessions import SecureCookieSessionInterface  # noqa: E402

from backend import captcha as captcha_lib  # noqa: E402
from backend.app import create_app  # noqa: E402
from backend.config import Config  # noqa: E402
from backend.db import init_db  # noqa: E402
from backend.security import reset_rate_limits  # noqa: E402
from backend.services import deploy as deploy_lib  # noqa: E402

failures = []


def check(name, ok, detail=""):
    status = "通过" if ok else "失败"
    print(f"[{status}] {name}" + (f"  ({detail})" if detail and not ok else ""))
    if not ok:
        failures.append(name)


def new_captcha(client):
    """取一张验证码并白盒读出答案：答案在服务端内存，token 在会话 cookie 里。"""
    resp = client.get("/api/auth/captcha")
    assert resp.status_code == 200
    cookie = client.get_cookie("session")
    assert cookie is not None, "验证码接口未写入会话"
    session_data = SecureCookieSessionInterface().get_signing_serializer(app).loads(cookie.value)
    token = session_data.get(captcha_lib.SESSION_KEY)
    entry = captcha_lib._entries.get(token)
    assert entry, "验证码答案未在服务端登记"
    return entry[0]


def try_login(client, username, password):
    return client.post(
        "/api/auth/login",
        json={"username": username, "password": password, "captcha": new_captcha(client)},
    )


def login(client, username, password):
    resp = try_login(client, username, password)
    assert resp.status_code == 200, f"登录失败 {username}: {resp.get_json()}"
    return resp


def component_id(name):
    row = next(item for item in client.get("/api/components").get_json()["data"] if item["name"] == name)
    return row["id"]


def open_position(component, expected_holder=None):
    """取某器件仍在外的持有点；可校验持有人。"""
    data = client.get(f"/api/custody/positions", query_string={"component_id": component}).get_json()["data"]
    if expected_holder is not None:
        row = next(item for item in data if item["holder"] == expected_holder)
    else:
        row = data[0]
    return row


init_db()
app = create_app()
client = app.test_client()
admin = client  # 主客户端始终以管理员身份操作

print("== 登录 ==")
resp = client.get("/api/auth/captcha")
check("验证码图片可获取", resp.status_code == 200 and resp.data[:4] == b"\x89PNG")
check("错误密码被拒绝", try_login(client, Config.USERNAME, "wrong-password").status_code == 401)
resp = login(client, Config.USERNAME, Config.PASSWORD)
check("总管理员登录成功", resp.get_json().get("role") == "superadmin")
resp = client.get("/api/auth/me")
check("会话信息含角色", resp.get_json()["user"] == Config.USERNAME and resp.get_json()["role"] == "superadmin")
check("健康检查", client.get("/api/health").status_code == 200)

print("== 器件管理 ==")
resp = client.post("/api/components", json={"name": "电阻10K", "asset_no": "ZC-001", "image_data": "data:image/png;base64,AA==", "category": "电阻", "spec": "10K 0805", "unit": "个", "location": "柜A-1", "owner": "张老师", "stock": 100, "threshold": 20})
check("新增器件", resp.status_code == 200)
rid = resp.get_json()["id"]
resp = client.post("/api/components", json={"name": "电容100uF", "category": "电容", "spec": "100uF 16V", "unit": "个", "location": "柜A-2", "stock": 50, "threshold": 10})
check("新增第二个器件", resp.status_code == 200)
check("重名器件被拒绝", client.post("/api/components", json={"name": "电阻10K", "spec": "10K 0805"}).status_code == 400)
check("固定资产编号不可重复", client.post("/api/components", json={"name": "资产编号重复件", "asset_no": "ZC-001"}).status_code == 400)
check("缺名称被拒绝", client.post("/api/components", json={"name": "  "}).status_code == 400)
check("初始库存超六位小数被拒绝", client.post("/api/components", json={"name": "精度坏例", "spec": "x", "stock": "0.1234567"}).status_code == 400)
check("初始库存负数被拒绝", client.post("/api/components", json={"name": "精度坏例", "spec": "x", "stock": -5}).status_code == 400)
check("器件列表", len(client.get("/api/components").get_json()["data"]) == 2)
check(
    "设备负责人已保存",
    any(item["owner"] == "张老师" for item in client.get("/api/components").get_json()["data"]),
)
first_row = next(item for item in client.get("/api/components").get_json()["data"] if item["id"] == rid)
check(
    "设备图片与固定资产编号已保存",
    first_row["asset_no"] == "ZC-001" and first_row["image_data"].startswith("data:image/png"),
)
uploaded_png = b"\x89PNG\r\n\x1a\n" + b"smoke-image"
resp = client.post(
    f"/api/components/{rid}/image",
    data={"image": (io.BytesIO(uploaded_png), "equipment.png")},
    content_type="multipart/form-data",
)
check("设备图片可独立上传", resp.status_code == 200)
first_row = next(item for item in client.get("/api/components").get_json()["data"] if item["id"] == rid)
check(
    "独立上传的设备图片已保存",
    first_row["image_data"].startswith("data:image/png;base64,"),
)
resp = client.put(f"/api/components/{rid}", json={"location": "柜B-1"})
check("编辑器件", resp.status_code == 200)
resp = client.put(f"/api/components/{rid}", json={"stock": 120})
adjusted_row = next(item for item in client.get("/api/components").get_json()["data"] if item["id"] == rid)
check("编辑设备数量", resp.status_code == 200 and adjusted_row["stock"] == 120)
adjust_record = client.get("/api/records", query_string={"cid": rid, "type": "adjust"}).get_json()
check(
    "设备数量修改写入调整流水",
    adjust_record["total"] == 1 and adjust_record["data"][0]["delta"] == 20,
)
check(
    "编辑设备数量不能为负数",
    client.put(f"/api/components/{rid}", json={"stock": -1}).status_code == 400,
)
client.put(f"/api/components/{rid}", json={"stock": 100})

resp = client.post(
    "/api/components",
    json={"name": "维修测试设备", "spec": "MAINT-1", "stock": 1, "maintenance": True},
)
maintenance_id = resp.get_json()["id"]
maintenance_row = next(
    item for item in client.get("/api/components").get_json()["data"]
    if item["id"] == maintenance_id
)
check("维修设备状态已保存", maintenance_row["equipment_status"] == "maintenance")
check(
    "维修设备不可借用",
    client.post(
        "/api/stock/out",
        json={"id": maintenance_id, "qty": 1, "applicant": "张三", "contact": "13800000000"},
    ).status_code == 400,
)
maintenance_stats = client.get("/api/stats/overview").get_json()["data"]
check("维修设备计入首页统计", maintenance_stats["maintenance_count"] == 1)
check("维修测试设备可清理", client.delete(f"/api/components/{maintenance_id}").status_code == 200)

resp = client.post(
    "/api/components",
    json={"name": "不可借测试设备", "spec": "BLOCK-1", "stock": 1, "unavailable": True},
)
unavailable_id = resp.get_json()["id"]
unavailable_row = next(
    item for item in client.get("/api/components").get_json()["data"]
    if item["id"] == unavailable_id
)
check("不可借状态已保存", unavailable_row["equipment_status"] == "unavailable")
check(
    "不可借设备不能办理借用",
    client.post(
        "/api/stock/out",
        json={"id": unavailable_id, "qty": 1, "applicant": "张三", "contact": "13800000000"},
    ).status_code == 400,
)
check("不可借测试设备可清理", client.delete(f"/api/components/{unavailable_id}").status_code == 200)

resp = client.post(
    "/api/components",
    json={"name": "混合状态测试设备", "spec": "MIXED-1", "stock": 10},
)
mixed_id = resp.get_json()["id"]
resp = client.put(
    f"/api/components/{mixed_id}",
    json={"total_qty": 10, "manual_status": "maintenance", "status_qty": 2},
)
mixed_row = next(item for item in client.get("/api/components").get_json()["data"] if item["id"] == mixed_id)
check(
    "部分维护设备仍显示可借用",
    resp.status_code == 200
    and mixed_row["equipment_status"] == "available"
    and mixed_row["stock"] == 8
    and mixed_row["maintenance_qty"] == 2,
)
mixed_stats = client.get("/api/stats/overview").get_json()["data"]
check(
    "主页按数量统计维护设备",
    mixed_stats["total_count"] == 160
    and mixed_stats["available_count"] == 158
    and mixed_stats["maintenance_count"] == 2,
)
resp = client.put(
    f"/api/components/{mixed_id}",
    json={"total_qty": 10, "manual_status": "unavailable", "status_qty": 3},
)
mixed_row = next(item for item in client.get("/api/components").get_json()["data"] if item["id"] == mixed_id)
check(
    "不可借数量可单独设置",
    resp.status_code == 200
    and mixed_row["stock"] == 7
    and mixed_row["maintenance_qty"] == 0
    and mixed_row["unavailable_qty"] == 3,
)
check("混合状态测试设备可清理", client.delete(f"/api/components/{mixed_id}").status_code == 200)

print("== 出入库与申请人 ==")
resp = client.post("/api/stock/in", json={"id": rid, "qty": 30})
check("入库成功", resp.status_code == 200 and abs(resp.get_json()["stock"] - 130) < 1e-9)
check("出库缺申请人被拒绝", client.post("/api/stock/out", json={"id": rid, "qty": 1}).status_code == 400)
check(
    "借用缺联系方式被拒绝",
    client.post("/api/stock/out", json={"id": rid, "qty": 1, "applicant": "张三"}).status_code == 400,
)
check(
    "邮箱不能作为借用联系方式",
    client.post("/api/stock/out", json={"id": rid, "qty": 1, "applicant": "张三", "contact": "zhang@example.com"}).status_code == 400,
)
resp = client.post(
    "/api/stock/out",
    json={"id": rid, "qty": 30, "applicant": "张三", "contact": "13800000000"},
)
check("出库成功并回写库存", resp.status_code == 200 and abs(resp.get_json()["stock"] - 100) < 1e-9)
borrow_chain = next(item for item in client.get("/api/custody", query_string={"q": "13800000000"}).get_json()["data"])
check(
    "借用联系方式与日期已保存",
    borrow_chain["contact"] == "13800000000"
    and borrow_chain["created_at"].startswith(date.today().isoformat()),
)
borrow_filter = client.get(
    "/api/custody", query_string={"q": "13800000000", "record_type": "borrow"}
).get_json()
return_filter = client.get(
    "/api/custody", query_string={"q": "13800000000", "record_type": "return"}
).get_json()
check("未归还借用链归入借用记录", borrow_filter["total"] == 1 and return_filter["total"] == 0)
borrowed_row = next(
    item for item in client.get("/api/components").get_json()["data"] if item["id"] == rid
)
check(
    "部分借出但仍有库存时状态为可借用",
    borrowed_row["equipment_status"] == "available"
    and abs(borrowed_row["borrowed_qty"] - 30) < 1e-9,
)
check("出库超量被拒绝", client.post("/api/stock/out", json={"id": rid, "qty": 999999, "applicant": "张三"}).status_code == 400)
check("出库超六位小数被拒绝", client.post("/api/stock/out", json={"id": rid, "qty": "0.1234567", "applicant": "张三"}).status_code == 400)
check("出库非数字被拒绝", client.post("/api/stock/out", json={"id": rid, "qty": "abc", "applicant": "张三"}).status_code == 400)
data = client.get("/api/records", query_string={"type": "out"}).get_json()
check("流水带申请人", data["data"][0]["applicant"] == "张三")
total_q = client.get("/api/records", query_string={"q": "张三"}).get_json()["total"]
check("按申请人检索流水", total_q >= 1)
data = client.get("/api/stats/overview").get_json()["data"]
check("库存汇总", abs(data["total_stock"] - 150) < 1e-6)
check(
    "首页按设备数量统计正确",
    data["kinds"] == 2
    and data["total_count"] == 180
    and data["available_count"] == 150
    and data["borrowed_count"] == 30
    and data["maintenance_count"] == 0
    and data["unavailable_count"] == 0,
)

print("== 流转精度（0.3 分次转交/归还） ==")
resp = client.post("/api/components", json={"name": "精度测试件", "spec": "P-0.3", "unit": "个", "stock": 0.3})
precision_id = resp.get_json()["id"]
resp = client.post("/api/stock/out", json={"id": precision_id, "qty": 0.3, "applicant": "王五", "contact": "13900000000"})
check("出库 0.3", resp.status_code == 200 and abs(resp.get_json()["stock"] - 0) < 1e-9)
pos = client.get("/api/custody/positions", query_string={"component_id": precision_id}).get_json()["data"]
check("流转链初始持有人", len(pos) == 1 and pos[0]["holder"] == "王五" and abs(pos[0]["qty"] - 0.3) < 1e-9)
root_id = pos[0]["root_id"]
check(
    "转借联系电话必填",
    client.post("/api/custody/transfer", json={"position_id": pos[0]["id"], "holder": "赵六", "qty": 0.1}).status_code == 400,
)
resp = client.post(
    "/api/custody/transfer",
    json={"position_id": pos[0]["id"], "holder": "赵六", "contact": "13700000001", "qty": 0.1},
)
check("部分转交 0.1", resp.status_code == 200)
resp = client.post(
    "/api/custody/transfer",
    json={"position_id": pos[0]["id"], "holder": "钱七", "contact": "13700000002", "qty": 0.2},
)
check("再次转交 0.2 不受浮点残量影响", resp.status_code == 200)
pos = client.get("/api/custody/positions", query_string={"component_id": precision_id}).get_json()["data"]
check("链上无浮点残量", abs(sum(item["qty"] for item in pos) - 0.3) < 1e-9 and len(pos) == 2)
zhao = next(item for item in pos if item["holder"] == "赵六")
qian = next(item for item in pos if item["holder"] == "钱七")
# 赵六只持有 0.1，转 0.2 必须被拒绝（修复前浮点残量会让这类判断失效）
check("超持有量转交被拒绝", client.post("/api/custody/transfer", json={"position_id": zhao["id"], "holder": "某人", "contact": "13700000003", "qty": 0.2}).status_code == 400)
check("超六位小数转交被拒绝", client.post("/api/custody/transfer", json={"position_id": zhao["id"], "holder": "某人", "qty": "0.1234567"}).status_code == 400)
check("零数量转交被拒绝", client.post("/api/custody/transfer", json={"position_id": zhao["id"], "holder": "某人", "qty": 0}).status_code == 400)
check("负数量转交被拒绝", client.post("/api/custody/transfer", json={"position_id": zhao["id"], "holder": "某人", "qty": -1}).status_code == 400)
check("超上限数量被拒绝", client.post("/api/custody/transfer", json={"position_id": zhao["id"], "holder": "某人", "qty": 1e12}).status_code == 400)
resp = client.post("/api/custody/return", json={"position_id": zhao["id"], "qty": 0.1})
first_return = resp.get_json()
check(
    "归还 0.1 回补库存并自动记录日期",
    resp.status_code == 200
    and abs(first_return["stock"] - 0.1) < 1e-9
    and first_return["returned_at"].startswith(date.today().isoformat()),
)
resp = client.post("/api/custody/return", json={"position_id": qian["id"], "qty": 0.2})
check("归还 0.2 库存精确到 0.3", resp.status_code == 200 and abs(resp.get_json()["stock"] - 0.3) < 1e-9)
check("超持有量归还被拒绝", client.post("/api/custody/return", json={"position_id": zhao["id"], "qty": 0.1}).status_code == 400)
returns_data = client.get("/api/records", query_string={"type": "return"}).get_json()
check("归还写入 return 类型流水", returns_data["total"] == 2)
detail = client.get(f"/api/custody/{root_id}/records").get_json()
kinds = [event["kind"] for event in detail["events"]]
check("链条完整轨迹", kinds == ["out", "transfer", "transfer", "return", "return"])
check(
    "借还记录显示最近归还时间",
    detail["summary"]["last_return_at"].startswith(date.today().isoformat())
    and all(event["created_at"] for event in detail["events"] if event["kind"] == "return"),
)
transfer_events = [event for event in detail["events"] if event["kind"] == "transfer"]
check(
    "转借记录保存联系方式和经手人",
    all(event["handler"] == "王五" for event in transfer_events)
    and [event["contact"] for event in transfer_events] == ["13700000001", "13700000002"],
)
holders = {item["holder"] for item in detail["positions"]}
check("历任持有人完整", holders == {"王五", "赵六", "钱七"})
q_data = client.get("/api/custody", query_string={"q": "王五"}).get_json()
check("已全部归还记录仍可查询", q_data["total"] == 1 and q_data["data"][0]["holding_total"] == 0)
q_data = client.get("/api/custody", query_string={"q": "P-0.3"}).get_json()
check("已归还设备可按规格检索", q_data["total"] == 1)
returned_filter = client.get(
    "/api/custody", query_string={"q": "P-0.3", "record_type": "return"}
).get_json()
borrowed_filter = client.get(
    "/api/custody", query_string={"q": "P-0.3", "record_type": "borrow"}
).get_json()
check("已归还借用链归入归还记录", returned_filter["total"] == 1 and borrowed_filter["total"] == 0)

print("== 三级用户与权限 ==")
student_account = "20260001"
student_initial_password = "swust0001"
teacher_account = "13800001111"
teacher_initial_password = "swust1111"
resp = client.post(
    "/api/users",
    json={"username": student_account, "full_name": "李同学", "phone": "13900000001", "user_type": "student", "mentor": "张老师"},
)
check("总管理员创建普通用户", resp.status_code == 201)
check("学生初始密码按学号后四位生成", resp.get_json()["initial_password"] == student_initial_password)
resp = client.post(
    "/api/users",
    json={"username": teacher_account, "full_name": "王老师", "user_type": "teacher", "role": "manager"},
)
check("总管理员创建管理员", resp.status_code == 201)
check("老师初始密码按手机号后四位生成", resp.get_json()["initial_password"] == teacher_initial_password)
check("所有用户姓名必填", client.post("/api/users", json={"username": "13900009999", "user_type": "teacher"}).status_code == 400)
check("老师账号必须是手机号", client.post("/api/users", json={"username": "1234", "full_name": "老师甲", "user_type": "teacher"}).status_code == 400)
check("学生账号必须是数字学号", client.post("/api/users", json={"username": "student-a", "full_name": "学生甲", "phone": "13900000002", "user_type": "student", "mentor": "张老师"}).status_code == 400)
check("学生所属导师必填", client.post("/api/users", json={"username": "20260009", "full_name": "学生乙", "phone": "13900000003", "user_type": "student"}).status_code == 400)
check("学生手机号必填", client.post("/api/users", json={"username": "20260010", "full_name": "学生丙", "user_type": "student", "mentor": "张老师"}).status_code == 400)
check("重名用户被拒绝", client.post("/api/users", json={"username": student_account, "full_name": "李同学", "phone": "13900000001", "user_type": "student", "mentor": "张老师"}).status_code == 400)
check("管理员名占用被拒绝", client.post("/api/users", json={"username": Config.USERNAME, "password": "tester123"}).status_code == 400)
users_list = client.get("/api/users").get_json()["data"]
check("用户列表含新用户", any(item["username"] == student_account for item in users_list))
tester_user_id = next(item["id"] for item in users_list if item["username"] == student_account)
operator_user_id = next(item["id"] for item in users_list if item["username"] == teacher_account)
check("用户列表含角色", next(item for item in users_list if item["id"] == operator_user_id)["role"] == "manager")
check(
    "用户列表区分老师和学生",
    next(item for item in users_list if item["id"] == tester_user_id)["user_type"] == "student"
    and next(item for item in users_list if item["id"] == tester_user_id)["full_name"] == "李同学"
    and next(item for item in users_list if item["id"] == tester_user_id)["phone"] == "13900000001"
    and next(item for item in users_list if item["id"] == tester_user_id)["mentor"] == "张老师"
    and next(item for item in users_list if item["id"] == operator_user_id)["user_type"] == "teacher",
)
resp = client.put(f"/api/users/{tester_user_id}/name", json={"full_name": "李明"})
updated_student = next(
    item for item in client.get("/api/users").get_json()["data"]
    if item["id"] == tester_user_id
)
check("用户姓名可修改并显示", resp.status_code == 200 and updated_student["full_name"] == "李明")
check(
    "用户姓名不能为空",
    client.put(f"/api/users/{tester_user_id}/name", json={"full_name": ""}).status_code == 400,
)
resp = client.put(f"/api/users/{tester_user_id}/mentor", json={"mentor": "李老师"})
updated_student = next(
    item for item in client.get("/api/users").get_json()["data"]
    if item["id"] == tester_user_id
)
check(
    "学生所属导师可补录并显示",
    resp.status_code == 200 and updated_student["mentor"] == "李老师",
)
check(
    "学生所属导师不能为空",
    client.put(f"/api/users/{tester_user_id}/mentor", json={"mentor": ""}).status_code == 400,
)
check(
    "老师账号不能设置所属导师",
    client.put(f"/api/users/{operator_user_id}/mentor", json={"mentor": "王老师"}).status_code == 400,
)
resp = client.put(f"/api/users/{tester_user_id}/phone", json={"phone": "13700000001"})
updated_student = next(
    item for item in client.get("/api/users").get_json()["data"]
    if item["id"] == tester_user_id
)
check("学生手机号可修改并显示", resp.status_code == 200 and updated_student["phone"] == "13700000001")

user_template = client.get("/api/users/template")
user_template_book = load_workbook(io.BytesIO(user_template.data), read_only=False)
user_template_sheet = user_template_book.active
user_template_headers = [cell.value for cell in next(user_template_sheet.iter_rows())]
user_validations = list(user_template_sheet.data_validations.dataValidation)
user_template_book.close()
check(
    "用户 Excel 模板包含姓名、所属导师及下拉选项",
    user_template.status_code == 200
    and user_template_headers == ["用户类型", "姓名", "账号", "手机号", "所属导师", "权限角色"]
    and len(user_validations) == 2,
)

user_book = Workbook()
user_sheet = user_book.active
user_sheet.append(["用户类型", "姓名", "账号", "手机号", "所属导师", "权限角色"])
user_sheet.append(["老师", "赵老师", "13900002222", "13900002222", "", "管理员"])
user_sheet.append(["学生", "周同学", "20260002", "13900002223", "李老师", "普通用户"])
user_sheet.append(["学生", "重复学生", student_account, "13900002224", "张老师", "普通用户"])
user_sheet.append(["未知", "未知用户", "20260003", "13900002225", "王老师", "普通用户"])
user_buffer = io.BytesIO()
user_book.save(user_buffer)
user_buffer.seek(0)
resp = client.post(
    "/api/users/import",
    data={"file": (user_buffer, "users.xlsx")},
    content_type="multipart/form-data",
)
import_result = resp.get_json()
check(
    "Excel 批量导入老师和学生",
    resp.status_code == 200 and import_result["imported"] == 2 and import_result["skipped"] == 2,
)
check(
    "导入用户使用自动初始密码",
    login(app.test_client(), "13900002222", "swust2222").status_code == 200
    and login(app.test_client(), "20260002", "swust0002").status_code == 200,
)

member = app.test_client()
resp = login(member, student_account, student_initial_password)
check("普通用户登录", resp.get_json().get("role") == "member")
check("普通用户可浏览", member.get("/api/components").status_code == 200)
check("普通用户新增器件被拒绝", member.post("/api/components", json={"name": "越权器件"}).status_code == 403)
check("普通用户编辑器件被拒绝", member.put(f"/api/components/{rid}", json={"location": "x"}).status_code == 403)
check("普通用户删除器件被拒绝", member.delete(f"/api/components/{rid}").status_code == 403)
check("普通用户批量导入被拒绝", member.post("/api/records/import").status_code == 403)
check("普通用户批量导入用户被拒绝", member.post("/api/users/import").status_code == 403)
check("普通用户管理账号被拒绝", member.get("/api/users").status_code == 403)
check("普通用户查询借还记录被拒绝", member.get("/api/custody").status_code == 403)
check("普通用户查询库存流水被拒绝", member.get("/api/records").status_code == 403)
check("普通用户入库被拒绝", member.post("/api/stock/in", json={"id": rid, "qty": 5}).status_code == 403)
check("普通用户借用被拒绝", member.post("/api/stock/out", json={"id": rid, "qty": 2, "applicant": "李四"}).status_code == 403)

manager = app.test_client()
resp = login(manager, teacher_account, teacher_initial_password)
check("管理员登录", resp.get_json().get("role") == "manager")
check("管理员不能管理账号", manager.get("/api/users").status_code == 403)
check("管理员不能编辑设备", manager.put(f"/api/components/{rid}", json={"location": "越权"}).status_code == 403)
check("管理员不能入库", manager.post("/api/stock/in", json={"id": rid, "qty": 5}).status_code == 403)
check("管理员不能查询库存流水", manager.get("/api/records").status_code == 403)
check("管理员可以查询借还记录", manager.get("/api/custody").status_code == 200)
manager_before = next(item for item in manager.get("/api/components").get_json()["data"] if item["id"] == rid)["stock"]
resp = manager.post("/api/stock/out", json={"id": rid, "qty": 2, "applicant": "李四", "contact": "0816-1234567"})
check("管理员可办理借用", resp.status_code == 200)
li_pos = open_position(rid, "李四")
check("出库自动建流转链", li_pos["holder"] == "李四" and abs(li_pos["qty"] - 2) < 1e-9)
check("普通用户转交被拒绝", member.post("/api/custody/transfer", json={"position_id": li_pos["id"], "holder": "越权", "qty": 1}).status_code == 403)
resp = manager.post("/api/custody/transfer", json={"position_id": li_pos["id"], "holder": "孙八", "contact": "13600000001", "qty": 1})
check("管理员可转交", resp.status_code == 200)
sun_pos = open_position(rid, "孙八")
check("普通用户归还被拒绝", member.post("/api/custody/return", json={"position_id": sun_pos["id"], "qty": 1}).status_code == 403)
resp = manager.post("/api/custody/return", json={"position_id": sun_pos["id"], "qty": 1})
check("管理员可归还", resp.status_code == 200 and abs(resp.get_json()["stock"] - (manager_before - 1)) < 1e-9)

resp = client.put(f"/api/users/{tester_user_id}/password", json={"password": "tester456"})
check("管理员重置密码", resp.status_code == 200)
check("重置后旧会话失效", member.get("/api/components").status_code == 401)
check("旧密码无法登录", try_login(member, student_account, student_initial_password).status_code == 401)
check("新密码可登录", login(member, student_account, "tester456").status_code == 200)
resp = client.delete(f"/api/users/{tester_user_id}")
check("管理员删除用户", resp.status_code == 200)
check("删除后会话失效", member.get("/api/components").status_code == 401)
check("已删除用户无法登录", try_login(member, student_account, "tester456").status_code == 401)
member = None
resp = client.delete(f"/api/users/{operator_user_id}")
check("总管理员删除管理员", resp.status_code == 200)
check("管理员删除后会话失效", manager.get("/api/components").status_code == 401)
manager = None

print("== 登录限流（IP+账号双维度）==")
# 用独立客户端：登录失败路径会 session.clear()，不能污染管理员的会话
brute = app.test_client()
for _ in range(Config.LOGIN_MAX_FAIL_PER_USER):
    assert try_login(brute, "locktest", "wrong-pw").status_code == 401
check("同账号超限后返回 429", try_login(brute, "locktest", "wrong-pw").status_code == 429)
check("账号锁定后正确密码也被拒", try_login(brute, "locktest", "any-pw").status_code == 429)
check("账号锁定不牵连其他账号", try_login(brute, Config.USERNAME, "still-wrong").status_code == 401)
for index in range(Config.LOGIN_MAX_FAIL):
    if try_login(brute, f"spray{index}", "wrong-pw").status_code == 429:
        break
check("同 IP 超限后任意账号返回 429", try_login(brute, Config.USERNAME, Config.PASSWORD).status_code == 429)
reset_rate_limits()
check("限流重置后可正常登录", try_login(brute, Config.USERNAME, Config.PASSWORD).status_code == 200)
brute = None

print("== 流转跨页检索 ==")
base_total = client.get("/api/custody").get_json()["total"]
for index in range(1, 13):
    resp = client.post("/api/components", json={"name": f"跨页器件{index}", "spec": f"K-{index}", "unit": "个", "stock": 1})
    cid = resp.get_json()["id"]
    resp = client.post("/api/stock/out", json={"id": cid, "qty": 1, "applicant": f"持有{index}", "contact": f"138000{index:05d}"})
    assert resp.status_code == 200, f"建链失败 {index}"
listed = client.get("/api/custody").get_json()
check("链条总数与分页一致", listed["total"] == base_total + 12 and listed["pages"] == -(-listed["total"] // 10))
page2 = client.get("/api/custody", query_string={"page": 2}).get_json()
check("第二页数量正确", len(page2["data"]) == listed["total"] - 10)
q_data = client.get("/api/custody", query_string={"q": "持有3"}).get_json()
check("后续页目标可被搜索直达", q_data["total"] == 1 and q_data["pages"] == 1)
holding2 = next(item for item in client.get("/api/custody", query_string={"q": "持有2"}).get_json()["data"])
resp = client.post("/api/custody/transfer", json={"position_id": open_position(holding2["component_id"], "持有2")["id"], "holder": "转交目标A", "contact": "13500000001", "qty": 0.5})
check("历史链可继续转交", resp.status_code == 200)
q_data = client.get("/api/custody", query_string={"q": "持有2"}).get_json()
check("按历史持有人检索", q_data["total"] == 1)
q_data = client.get("/api/custody", query_string={"q": "转交目标A"}).get_json()
check("按新持有人检索", q_data["total"] == 1 and "转交目标A" in [h["holder"] for h in q_data["data"][0]["holdings"]])
check("按规格检索", client.get("/api/custody", query_string={"q": "K-5"}).get_json()["total"] == 1)
check("无关关键词返回空", client.get("/api/custody", query_string={"q": "不存在xyz"}).get_json()["total"] == 0)

print("== 预计归还日期与逾期 ==")
future = (date.today() + timedelta(days=7)).isoformat()
resp = client.post("/api/components", json={"name": "逾期测试件", "spec": "DUE-1", "unit": "个", "stock": 5})
due_id = resp.get_json()["id"]
client.post("/api/stock/out", json={"id": due_id, "qty": 2, "applicant": "日期持有A", "contact": "010-12345678", "due_date": future})
chain = next(item for item in client.get("/api/custody", query_string={"q": "DUE-1"}).get_json()["data"])
check("出库可带预计归还日期", chain["due_date"] == future and chain["overdue_days"] is None)
due_pos = open_position(due_id, "日期持有A")
client.post("/api/custody/transfer", json={"position_id": due_pos["id"], "holder": "日期持有B", "contact": "13400000001", "qty": 1})
detail = client.get(f"/api/custody/{chain['id']}/records").get_json()
child = next(e for e in detail["events"] if e["kind"] == "transfer")
check("转交自动继承归还日期", child["due_date"] == future)
new_due = (date.today() + timedelta(days=30)).isoformat()
client.post("/api/custody/transfer", json={"position_id": due_pos["id"], "holder": "日期持有C", "contact": "13400000002", "qty": 1, "due_date": new_due})
detail = client.get(f"/api/custody/{chain['id']}/records").get_json()
check("转交可显式改约新日期", [e for e in detail["events"] if e["kind"] == "transfer"][-1]["due_date"] == new_due)
check("过去日期被拒绝", client.post("/api/stock/out", json={"id": due_id, "qty": 1, "applicant": "日期持有D", "contact": "010-87654321", "due_date": "2020-01-01"}).status_code == 400)
check("乱格式日期被拒绝", client.post("/api/stock/out", json={"id": due_id, "qty": 1, "applicant": "日期持有D", "contact": "010-87654321", "due_date": "明天"}).status_code == 400)
# 制造逾期：白盒把在外持有点的日期改到三天前，验证逾期标记、筛选与统计
raw = sqlite3.connect(TMP_DB)
raw.execute(
    "UPDATE custody_positions SET due_date = ? WHERE component_id = ? AND ROUND(qty, 6) > 0",
    ((date.today() - timedelta(days=3)).isoformat(), due_id),
)
raw.commit()
raw.close()
chain = next(item for item in client.get("/api/custody", query_string={"q": "DUE-1"}).get_json()["data"])
check("逾期天数正确计算", chain["overdue_days"] == 3)
stats = client.get("/api/stats/overview").get_json()["data"]
check("仪表盘统计逾期链条数", stats["overdue_count"] == 1)

print("== 删除保护 ==")
check("有流转历史的器件禁止删除", client.delete(f"/api/components/{precision_id}").status_code == 400)
check("电阻10K（已有流转链）禁止删除", client.delete(f"/api/components/{rid}").status_code == 400)
resp = client.post("/api/components", json={"name": "待删除件", "spec": "tmp", "unit": "个", "stock": 7})
tmp_id = resp.get_json()["id"]
resp = client.delete(f"/api/components/{tmp_id}")
check("无流转器件可删除", resp.status_code == 200)

print("== 导入导出 ==")
template = client.get("/api/records/template")
template_book = load_workbook(io.BytesIO(template.data), read_only=True)
template_headers = [cell.value for cell in next(template_book.active.iter_rows())]
template_book.close()
check(
    "Excel 模板下载",
    template.status_code == 200
    and "spreadsheetml" in template.content_type
    and template_headers == ["设备名称", "设备固定资产编号", "分类", "设备型号", "单位", "位置", "设备负责人", "设备状态", "设备数量", "预警值", "备注"],
)
exported = client.get("/api/records/export")
check("清单导出", exported.status_code == 200 and "固定资产编号".encode("utf-8") in exported.data)
csv_body = (
    "名称,分类,规格,单位,位置,数量,阈值,备注\n"
    "导入测试件,测试,IMP-1,个,柜C-1,0.123456789,0,浮点尾巴归整\n"
    "导入第二件,测试,IMP-2,个,柜C-2,10,0,\n"
)
resp = client.post(
    "/api/records/import",
    data={"file": (io.BytesIO(csv_body.encode("utf-8")), "import.csv", "text/csv")},
    content_type="multipart/form-data",
)
check("批量导入成功", resp.status_code == 200 and resp.get_json()["imported"] == 2)
imp_stock = next(item["stock"] for item in client.get("/api/components", query_string={"q": "IMP-1"}).get_json()["data"])
check("导入数量按六位归整", abs(imp_stock - 0.123457) < 1e-9)

xlsx_book = Workbook()
xlsx_sheet = xlsx_book.active
xlsx_sheet.append(["设备名称", "设备固定资产编号", "分类", "设备型号", "单位", "位置", "设备负责人", "设备状态", "设备数量", "预警值", "备注"])
xlsx_sheet.append(["Excel导入设备", "XLSX-001", "机器人", "XR-1", "台", "实验室A区", "王老师", "可借用", 3, 1, "Excel测试"])
xlsx_file = io.BytesIO()
xlsx_book.save(xlsx_file)
xlsx_book.close()
xlsx_file.seek(0)
resp = client.post(
    "/api/records/import",
    data={"file": (xlsx_file, "equipment.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    content_type="multipart/form-data",
)
xlsx_row = next(item for item in client.get("/api/components", query_string={"q": "XLSX-001"}).get_json()["data"])
check(
    "Excel 设备信息导入",
    resp.status_code == 200 and xlsx_row["stock"] == 3 and xlsx_row["owner"] == "王老师",
)

print("== 自动部署 Webhook ==")
with open(SMOKE_SCRIPT, "w", encoding="utf-8") as fh:
    fh.write(f"#!/usr/bin/env bash\necho ok > '{SMOKE_MARKER}'\n")
payload = b'{"ref": "refs/heads/master"}'
bad = client.post("/webhook", data=payload, headers={"X-Hub-Signature-256": "sha256=deadbeef", "X-GitHub-Event": "push"})
check("非法签名被拒绝", bad.status_code == 403)
ignored = client.post("/webhook", data=payload, headers={"X-Hub-Signature-256": deploy_lib.sign(SMOKE_SECRET, payload), "X-GitHub-Event": "ping"})
check("非 push 事件忽略", ignored.status_code == 200)
if shutil.which("bash"):
    signed = client.post("/webhook", data=payload, headers={"X-Hub-Signature-256": deploy_lib.sign(SMOKE_SECRET, payload), "X-GitHub-Event": "push"})
    check("合法签名触发部署", signed.status_code == 200)
    marker_ok = any(os.path.exists(SMOKE_MARKER) for _ in range(40) if not time.sleep(0.1))
    check("部署脚本被执行", marker_ok)
else:
    print("[跳过] 合法签名触发部署与脚本执行（当前系统没有 bash）")

print("== 前端静态资源 ==")
home = client.get("/")
check("SPA 首页", home.status_code == 200 and b'id="app"' in home.data)
deep = client.get("/custody")
check("前端路由深度访问", deep.status_code == 200 and b'id="app"' in deep.data)
check("缺失静态资源 404", client.get("/assets/does-not-exist.js").status_code == 404)
assets_dir = os.path.join(Config.STATIC_DIST, "assets")
if os.path.isdir(assets_dir):
    asset = next((name for name in os.listdir(assets_dir) if name.endswith(".js")), None)
    if asset:
        cached = client.get(f"/assets/{asset}")
        check("哈希资源长期强缓存", "immutable" in (cached.headers.get("Cache-Control") or ""))

for suffix in ("", "-wal", "-shm"):
    if os.path.exists(TMP_DB + suffix):
        os.remove(TMP_DB + suffix)

for path in (SMOKE_SCRIPT, SMOKE_MARKER, SMOKE_LOG):
    if os.path.exists(path):
        os.remove(path)

print()
if failures:
    print(f"失败 {len(failures)} 项: {failures}")
    sys.exit(1)
print("全部通过")
