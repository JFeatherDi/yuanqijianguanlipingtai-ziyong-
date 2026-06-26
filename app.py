"""实验室元器件管理平台 - 后端"""
import os
import io
import csv
import sqlite3
from datetime import datetime
from functools import wraps

from flask import (
    Flask, request, session, jsonify, send_from_directory,
    Response, g
)
from openpyxl import load_workbook

# ---------- 配置 ----------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "data.db")
STATIC_DIR = os.path.join(BASE_DIR, "static")

USERNAME = "IOTAT"
PASSWORD = "swust350351"
SECRET_KEY = "swust-lab-components-2026"

app = Flask(__name__, static_folder=None)
app.secret_key = SECRET_KEY
app.config["MAX_CONTENT_LENGTH"] = 8 * 1024 * 1024  # 8MB 上传上限


# ---------- 数据库 ----------
def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
    return g.db


@app.teardown_appcontext
def close_db(exc):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db():
    conn = sqlite3.connect(DB_PATH)
    conn.executescript("""
    CREATE TABLE IF NOT EXISTS components (
        id          INTEGER PRIMARY KEY AUTOINCREMENT,
        name        TEXT NOT NULL,
        category    TEXT,
        spec        TEXT,
        unit        TEXT,
        location    TEXT,
        stock       REAL NOT NULL DEFAULT 0,
        threshold   REAL DEFAULT 0,
        remark      TEXT,
        UNIQUE(name, spec)
    );
    CREATE TABLE IF NOT EXISTS transactions (
        id          INTEGER PRIMARY KEY AUTOINCREMENT,
        component_id INTEGER NOT NULL,
        delta       REAL NOT NULL,
        type        TEXT NOT NULL,   -- in / out / import / init
        operator    TEXT,
        remark      TEXT,
        created_at  TEXT NOT NULL,
        FOREIGN KEY(component_id) REFERENCES components(id)
    );
    CREATE INDEX IF NOT EXISTS idx_comp_name ON components(name);
    CREATE INDEX IF NOT EXISTS idx_txn_comp ON transactions(component_id);
    """)
    conn.commit()
    conn.close()


# ---------- 鉴权 ----------
def login_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if not session.get("user"):
            return jsonify({"ok": False, "msg": "未登录或会话过期"}), 401
        return f(*args, **kwargs)
    return wrapper


# ---------- 工具 ----------
def now_str():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def row_to_dict(row):
    return {k: row[k] for k in row.keys()} if row else None


def apply_change(db, comp_id, delta, ttype, operator, remark=""):
    """统一处理库存变更 + 流水写入。delta 负数=出库。"""
    db.execute(
        "UPDATE components SET stock = stock + ? WHERE id = ?",
        (delta, comp_id),
    )
    db.execute(
        "INSERT INTO transactions (component_id, delta, type, operator, remark, created_at) "
        "VALUES (?, ?, ?, ?, ?, ?)",
        (comp_id, delta, ttype, operator, remark, now_str()),
    )


# ---------- 页面 ----------
@app.route("/")
def index():
    if not session.get("user"):
        return send_from_directory(STATIC_DIR, "login.html")
    return send_from_directory(STATIC_DIR, "index.html")


@app.route("/login", methods=["POST"])
def login():
    data = request.get_json(silent=True) or {}
    u = (data.get("username") or "").strip()
    p = (data.get("password") or "").strip()
    if u == USERNAME and p == PASSWORD:
        session["user"] = u
        return jsonify({"ok": True, "user": u})
    return jsonify({"ok": False, "msg": "账号或密码错误"}), 401


@app.route("/logout", methods=["POST"])
def logout():
    session.clear()
    return jsonify({"ok": True})


@app.route("/me")
@login_required
def me():
    return jsonify({"ok": True, "user": session.get("user")})


# ---------- 元器件 CRUD ----------
@app.route("/api/components")
@login_required
def list_components():
    db = get_db()
    keyword = (request.args.get("q") or "").strip()
    sql = "SELECT * FROM components"
    args = []
    if keyword:
        sql += " WHERE name LIKE ? OR category LIKE ? OR spec LIKE ? OR location LIKE ?"
        like = f"%{keyword}%"
        args = [like, like, like, like]
    sql += " ORDER BY category, name"
    rows = db.execute(sql, args).fetchall()
    return jsonify({"ok": True, "data": [row_to_dict(r) for r in rows]})


@app.route("/api/components", methods=["POST"])
@login_required
def add_component():
    db = get_db()
    d = request.get_json(silent=True) or {}
    init_stock = float(d.get("stock") or 0)
    try:
        cur = db.execute(
            "INSERT INTO components (name, category, spec, unit, location, stock, threshold, remark) "
            "VALUES (?, ?, ?, ?, ?, 0, ?, ?)",
            (
                d.get("name", "").strip(),
                d.get("category", "").strip(),
                d.get("spec", "").strip(),
                d.get("unit", "").strip() or "个",
                d.get("location", "").strip(),
                float(d.get("threshold") or 0),
                d.get("remark", "").strip(),
            ),
        )
        comp_id = cur.lastrowid
        if init_stock != 0:
            apply_change(db, comp_id, init_stock, "init",
                         session.get("user"), "新建初始库存")
        db.commit()
        return jsonify({"ok": True, "id": comp_id})
    except sqlite3.IntegrityError:
        return jsonify({"ok": False, "msg": "器件名称+规格已存在"}), 400
    except Exception as e:
        return jsonify({"ok": False, "msg": str(e)}), 400


@app.route("/api/components/<int:cid>", methods=["PUT"])
@login_required
def update_component(cid):
    db = get_db()
    d = request.get_json(silent=True) or {}
    fields = ("name", "category", "spec", "unit", "location", "threshold", "remark")
    sets = []
    args = []
    for f in fields:
        if f in d:
            sets.append(f"{f} = ?")
            args.append(d[f])
    if sets:
        args.append(cid)
        db.execute(f"UPDATE components SET {', '.join(sets)} WHERE id = ?", args)
        db.commit()
    return jsonify({"ok": True})


@app.route("/api/components/<int:cid>", methods=["DELETE"])
@login_required
def delete_component(cid):
    db = get_db()
    db.execute("DELETE FROM transactions WHERE component_id = ?", (cid,))
    db.execute("DELETE FROM components WHERE id = ?", (cid,))
    db.commit()
    return jsonify({"ok": True})


# ---------- 入库 / 出库 ----------
@app.route("/api/stock/in", methods=["POST"])
@login_required
def stock_in():
    db = get_db()
    d = request.get_json(silent=True) or {}
    cid = d.get("id")
    qty = float(d.get("qty") or 0)
    if qty <= 0:
        return jsonify({"ok": False, "msg": "数量必须大于 0"}), 400
    comp = db.execute("SELECT * FROM components WHERE id = ?", (cid,)).fetchone()
    if not comp:
        return jsonify({"ok": False, "msg": "器件不存在"}), 404
    apply_change(db, cid, qty, "in", session.get("user"), d.get("remark", ""))
    db.commit()
    return jsonify({"ok": True, "stock": comp["stock"] + qty})


@app.route("/api/stock/out", methods=["POST"])
@login_required
def stock_out():
    db = get_db()
    d = request.get_json(silent=True) or {}
    cid = d.get("id")
    qty = float(d.get("qty") or 0)
    if qty <= 0:
        return jsonify({"ok": False, "msg": "数量必须大于 0"}), 400
    comp = db.execute("SELECT * FROM components WHERE id = ?", (cid,)).fetchone()
    if not comp:
        return jsonify({"ok": False, "msg": "器件不存在"}), 404
    if comp["stock"] < qty:
        return jsonify({"ok": False, "msg": f"库存不足，当前 {comp['stock']}"}), 400
    apply_change(db, cid, -qty, "out", session.get("user"), d.get("remark", ""))
    db.commit()
    return jsonify({"ok": True, "stock": comp["stock"] - qty})


# ---------- 导入 ----------
@app.route("/api/import", methods=["POST"])
@login_required
def import_file():
    db = get_db()
    if "file" not in request.files:
        return jsonify({"ok": False, "msg": "未上传文件"}), 400
    f = request.files["file"]
    filename = (f.filename or "").lower()

    try:
        raw = f.read()
        if filename.endswith(".csv"):
            # 尝试 utf-8-sig，失败回退 gbk（Excel 中文 CSV 常为 gbk）
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = raw.decode("gbk", errors="replace")
            reader = csv.reader(io.StringIO(text))
            rows = list(reader)
        elif filename.endswith((".xlsx", ".xlsm")):
            wb = load_workbook(filename=io.BytesIO(raw), read_only=True, data_only=True)
            ws = wb.active
            rows = [[c.value for c in r] for r in ws.iter_rows()]
        else:
            return jsonify({"ok": False, "msg": "仅支持 .csv / .xlsx"}), 400

        if not rows:
            return jsonify({"ok": False, "msg": "文件为空"}), 400

        # 识别表头；找不到则默认第一列为 name
        header = [str(x).strip().lower() if x is not None else "" for x in rows[0]]
        col_map = {}
        for idx, h in enumerate(header):
            if h in ("name", "名称", "器件", "器件名称"):
                col_map["name"] = idx
            elif h in ("category", "分类", "类别"):
                col_map["category"] = idx
            elif h in ("spec", "规格", "型号"):
                col_map["spec"] = idx
            elif h in ("unit", "单位"):
                col_map["unit"] = idx
            elif h in ("location", "位置", "存放位置"):
                col_map["location"] = idx
            elif h in ("stock", "数量", "库存", "当前数量"):
                col_map["stock"] = idx
            elif h in ("threshold", "阈值", "预警值"):
                col_map["threshold"] = idx
            elif h in ("remark", "备注"):
                col_map["remark"] = idx

        if "name" not in col_map:
            # 第一种情况：第一列强制作为 name
            col_map["name"] = 0

        data_rows = rows[1:] if any(col_map.values()) else rows
        ok = 0
        skip = 0
        for r in data_rows:
            if not r or all((v is None or str(v).strip() == "") for v in r):
                continue
            def pick(key, default=""):
                idx = col_map.get(key)
                if idx is None or idx >= len(r):
                    return default
                v = r[idx]
                return "" if v is None else str(v).strip()

            name = pick("name")
            if not name:
                skip += 1
                continue
            spec = pick("spec")
            stock_raw = pick("stock", "0")
            try:
                stock = float(stock_raw) if stock_raw else 0.0
            except ValueError:
                stock = 0.0
            try:
                threshold = float(pick("threshold", "0")) or 0.0
            except ValueError:
                threshold = 0.0

            existing = db.execute(
                "SELECT id, stock FROM components WHERE name = ? AND IFNULL(spec,'') = IFNULL(?, '')",
                (name, spec),
            ).fetchone()
            if existing:
                if stock:
                    apply_change(db, existing["id"], stock, "import",
                                 session.get("user"), "导入累加")
                ok += 1
            else:
                cur = db.execute(
                    "INSERT INTO components (name, category, spec, unit, location, stock, threshold, remark) "
                    "VALUES (?, ?, ?, ?, ?, 0, ?, ?)",
                    (name, pick("category"), spec, pick("unit", "个"),
                     pick("location"), threshold, pick("remark")),
                )
                if stock:
                    apply_change(db, cur.lastrowid, stock, "init",
                                 session.get("user"), "导入初始")
                ok += 1
        db.commit()
        return jsonify({"ok": True, "imported": ok, "skipped": skip})
    except Exception as e:
        db.rollback()
        return jsonify({"ok": False, "msg": f"解析失败: {e}"}), 500


# ---------- 流水 ----------
@app.route("/api/transactions")
@login_required
def list_transactions():
    db = get_db()
    cid = request.args.get("cid")
    limit = min(int(request.args.get("limit", 100)), 500)
    if cid:
        rows = db.execute(
            "SELECT t.*, c.name, c.spec FROM transactions t "
            "JOIN components c ON c.id = t.component_id "
            "WHERE t.component_id = ? ORDER BY t.id DESC LIMIT ?",
            (cid, limit),
        ).fetchall()
    else:
        rows = db.execute(
            "SELECT t.*, c.name, c.spec FROM transactions t "
            "JOIN components c ON c.id = t.component_id "
            "ORDER BY t.id DESC LIMIT ?",
            (limit,),
        ).fetchall()
    return jsonify({"ok": True, "data": [row_to_dict(r) for r in rows]})


# ---------- 导出 ----------
@app.route("/api/export")
@login_required
def export_csv():
    db = get_db()
    rows = db.execute(
        "SELECT name, category, spec, unit, location, stock, threshold, remark "
        "FROM components ORDER BY category, name"
    ).fetchall()
    out = io.StringIO()
    out.write("\ufeff")  # BOM for Excel
    w = csv.writer(out)
    w.writerow(["名称", "分类", "规格", "单位", "位置", "库存", "预警值", "备注"])
    for r in rows:
        w.writerow([r["name"], r["category"], r["spec"], r["unit"],
                    r["location"], r["stock"], r["threshold"], r["remark"]])
    return Response(
        out.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=components.csv"},
    )


# ---------- 模板下载 ----------
@app.route("/api/template")
@login_required
def template():
    csv_text = "名称,分类,规格,单位,位置,库存,预警值,备注\n电阻10K,电阻,10K 0805,个,柜A-1,100,20,示例\n"
    return Response(
        "\ufeff" + csv_text,
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=template.csv"},
    )


if __name__ == "__main__":
    init_db()
    app.run(host="0.0.0.0", port=5000, threaded=True)